import json
import os
from io import BytesIO
from tempfile import NamedTemporaryFile
from unittest.mock import MagicMock

import pytest
from docx import Document
from fpdf import FPDF
from pytest_mock import MockFixture

FILE_CONTENT: str = "This is a sample document."


@pytest.fixture
def generate_pdf_file():
    with NamedTemporaryFile(suffix=".pdf", mode="wb", delete=False) as temp_file:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 100, txt=FILE_CONTENT, align="C")
        pdf.output(temp_file.name)
    yield temp_file.name

    os.remove(temp_file.name)


@pytest.fixture
def generate_doc_file():
    with NamedTemporaryFile(suffix=".docx", mode="wb", delete=False) as temp_file:
        doc = Document()
        doc.add_paragraph(FILE_CONTENT)
        doc.save(temp_file.name)

    yield temp_file.name

    os.remove(temp_file.name)

@pytest.fixture
def mock_invoke_model(mocker: MockFixture) -> MagicMock:
    sample_response = {"generation": '{"applicant_name": "Ray"}'}

    mock_response_get = BytesIO(json.dumps(sample_response).encode("utf-8"))

    invoke_model_response = mocker.MagicMock()
    invoke_model_response.get.return_value = mock_response_get
    invoke_model_mock_version = mocker.MagicMock()
    boto3_response = mocker.MagicMock()

    boto3_response.invoke_model = invoke_model_mock_version
    boto3_response.invoke_model.return_value = invoke_model_response

    mock_boto3_client = mocker.patch("boto3.client")
    mock_boto3_client.return_value = boto3_response

    return invoke_model_mock_version